"""Lecteurs de formats : tout document est ramené à des « grilles » (tableaux de cellules texte) dont chaque cellule
garde son emplacement (feuille + cellule A1, tableau + ligne, ligne de texte...). L'extraction des lignes d'offre et des
totaux est ensuite la même que pour un PDF : mêmes règles, mêmes références vers les deux documents.

Formats : PDF (voir extract.py), Excel .xlsx, Word .docx, OpenDocument .ods/.odt, CSV/TSV, HTML, JSON, texte/Markdown.
Hors périmètre, refusés explicitement : images et scans (pas d'OCR), anciens formats binaires .xls/.doc, formats inconnus.
Aucun appel à un modèle : lecture déterministe.
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal
from html.parser import HTMLParser
from typing import Any, List, Optional, Tuple
from xml.etree import ElementTree as ET

from .parsing import clean_spaces

MAX_ROWS, MAX_COLS, MAX_GRIDS = 2000, 40, 30


class UnsupportedFormat(ValueError):
    """Format que l'outil ne sait pas (ou ne veut pas) lire ; le message est montré à l'utilisateur."""


@dataclass
class Grid:
    name: str                 # « Offre » (feuille), « Tableau 2 », « Texte », « Fichier »...
    style: str                # a1 | table | line | item | field : comment l'emplacement est exprimé
    rows: List[List[str]]
    ref_prefix: str = ""      # style a1 : nom de feuille


@dataclass
class Loaded:
    kind: str
    grids: List[Grid]
    hints: str = ""           # texte additionnel pour la détection de devise (formats de cellules Excel)
    warnings: List[str] = field(default_factory=list)


def col_letter(i: int) -> str:
    s, i = "", i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _norm_rows(rows: List[List[str]]) -> List[List[str]]:
    rows = [[clean_spaces(str(c)) if c is not None else "" for c in r][:MAX_COLS] for r in rows[:MAX_ROWS]]
    while rows and not any(rows[-1]):
        rows.pop()
    w = max((len(r) for r in rows), default=0)
    return [r + [""] * (w - len(r)) for r in rows]


def _num_text(v: Any) -> str:
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, (float, Decimal)):
        d = Decimal(repr(v)) if isinstance(v, float) else v
        if d == d.to_integral():
            return str(int(d))
        return format(d.normalize(), "f")
    if isinstance(v, _dt.datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, _dt.date):
        return v.isoformat()
    return "" if v is None else str(v)


# --------------------------------------------------------------------------------------------------- détection
def detect_kind(data: bytes, filename: str = "") -> str:
    ext = (filename.rsplit(".", 1)[-1].lower() if "." in filename else "")
    head = data[:8]
    if data.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        try:
            zf = zipfile.ZipFile(io.BytesIO(data))
            names = set(zf.namelist())
            if sum(i.file_size for i in zf.infolist()) > 60 * 1024 * 1024:
                raise UnsupportedFormat("Fichier trop volumineux une fois décompressé.")
        except zipfile.BadZipFile:
            raise UnsupportedFormat("Archive illisible ou corrompue.")
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "content.xml" in names:
            mt = zipfile.ZipFile(io.BytesIO(data)).read("mimetype").decode("ascii", "ignore") if "mimetype" in names else ""
            if "spreadsheet" in mt or ext == "ods":
                return "ods"
            if "text" in mt or ext == "odt":
                return "odt"
        raise UnsupportedFormat("Archive non reconnue (ni Word, ni Excel, ni OpenDocument).")
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        raise UnsupportedFormat("Ancien format binaire Office (.xls / .doc) non pris en charge : enregistrer le fichier en .xlsx ou .docx.")
    if head.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"II*\x00", b"MM\x00*", b"BM")) or data[:12].startswith(b"RIFF") and b"WEBP" in data[:16]:
        return "image"
    if head.startswith(b"{\\rtf"):
        raise UnsupportedFormat("Format RTF non pris en charge : enregistrer le fichier en .docx ou en PDF.")
    if b"\x00" in data[:2048]:
        raise UnsupportedFormat("Format binaire non reconnu.")
    text = _decode(data)
    s = text.lstrip()
    if ext in ("html", "htm") or re.match(r"(?is)(<!doctype html|<html|<table|<body|<div|<p[ >])", s):
        return "html"
    if ext == "json" or (s[:1] in "{[" and _try_json(s) is not None):
        return "json"
    if ext in ("csv", "tsv"):
        return "csv"
    return "txt"


def _try_json(s: str):
    try:
        return json.loads(s)
    except ValueError:
        return None


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", "replace")


# --------------------------------------------------------------------------------------------------- Excel
def _read_xlsx(data: bytes) -> Loaded:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        wbf = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    except Exception:
        raise UnsupportedFormat("Classeur Excel illisible (fichier corrompu ou protégé par mot de passe).")
    out = Loaded("xlsx", [])
    missing = 0
    hidden = []
    cur_hints: List[str] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            hidden.append(ws.title)
            continue
        wsf = wbf[ws.title]
        rows: List[List[str]] = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, MAX_ROWS), max_col=min(ws.max_column or 0, MAX_COLS)):
            cells = []
            for c in row:
                v = c.value
                if v is None and isinstance(wsf[c.coordinate].value, str) and wsf[c.coordinate].value.startswith("="):
                    missing += 1
                if isinstance(v, (int, float, Decimal)) and c.number_format:
                    cur_hints.append(c.number_format)
                cells.append(_num_text(v))
            rows.append(cells)
        rows = _norm_rows(rows)
        if rows:
            out.grids.append(Grid(ws.title, "a1", rows, ref_prefix=ws.title))
    out.hints = " ".join(cur_hints[:200])
    if hidden:
        out.warnings.append(f"Feuilles masquées ignorées : {hidden}.")
    if missing:
        out.warnings.append(f"{missing} cellule(s) contiennent une formule sans valeur enregistrée (fichier jamais ouvert dans Excel) : "
                            "ignorées, rien n'est recalculé ni deviné.")
    return out


# --------------------------------------------------------------------------------------------------- Word
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _w_text(el) -> str:
    parts = []
    for n in el.iter():
        if n.tag == W + "t":
            parts.append(n.text or "")
        elif n.tag in (W + "tab", W + "br"):
            parts.append(" ")
    return clean_spaces("".join(parts))


def _w_blocks(parent):
    for ch in parent:
        if ch.tag == W + "sdt":
            c = ch.find(W + "sdtContent")
            if c is not None:
                yield from _w_blocks(c)
        else:
            yield ch


def _read_docx(data: bytes) -> Loaded:
    try:
        root = ET.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    except Exception:
        raise UnsupportedFormat("Document Word illisible.")
    body = root.find(W + "body")
    tables: List[Grid] = []
    paras: List[str] = []
    for blk in _w_blocks(body if body is not None else []):
        if blk.tag == W + "p":
            t = _w_text(blk)
            if t:
                paras.append(t)
        elif blk.tag == W + "tbl" and len(tables) < MAX_GRIDS:
            rows = []
            for tr in blk.findall(W + "tr"):
                cells: List[str] = []
                for tc in tr.findall(W + "tc"):
                    span = tc.find(f"{W}tcPr/{W}gridSpan")
                    n = int(span.get(W + "val", "1")) if span is not None else 1
                    cells.append(" ".join(x for x in (_w_text(p) for p in tc.iter(W + "p")) if x))
                    cells.extend([""] * (n - 1))
                rows.append(cells)
            rows = _norm_rows(rows)
            if rows:
                tables.append(Grid(f"Tableau {len(tables) + 1}", "table", rows))
    grids = tables + ([Grid("Texte", "line", _norm_rows([[p] for p in paras]))] if paras else [])
    return Loaded("docx", grids)


# --------------------------------------------------------------------------------------------------- OpenDocument
def _read_odf(data: bytes, kind: str) -> Loaded:
    T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
    O = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
    X = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    try:
        root = ET.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("content.xml"))
    except Exception:
        raise UnsupportedFormat("Document OpenDocument illisible.")

    def cell_text(c) -> str:
        vt = c.get(O + "value-type")
        if vt in ("float", "percentage", "currency") and c.get(O + "value") is not None:
            return _num_text(Decimal(c.get(O + "value")))
        if vt == "date" and c.get(O + "date-value"):
            return c.get(O + "date-value")[:10]
        return clean_spaces(" ".join("".join(p.itertext()) for p in c.findall(X + "p")))

    grids: List[Grid] = []
    for tb in root.iter(T + "table"):
        rows = []
        for tr in tb.iter(T + "table-row"):
            cells: List[str] = []
            for c in tr:
                if c.tag not in (T + "table-cell", T + "covered-table-cell"):
                    continue
                rep = min(int(c.get(T + "number-columns-repeated", "1")), MAX_COLS)
                cells.extend([cell_text(c)] * rep)
            rows.append(cells)
        rows = _norm_rows(rows)
        name = tb.get(T + "name") or f"Tableau {len(grids) + 1}"
        if rows and len(grids) < MAX_GRIDS:
            grids.append(Grid(name, "a1" if kind == "ods" else "table", rows, ref_prefix=name))
    if kind == "odt":
        # les paragraphes de cellules de tableau sont aussi des text:p : on ne garde que ceux hors tableau
        in_tables = set()
        for tb in root.iter(T + "table"):
            for p in tb.iter(X + "p"):
                in_tables.add(id(p))
        paras = [clean_spaces("".join(p.itertext())) for p in root.iter(X + "p") if id(p) not in in_tables]
        paras = [p for p in paras if p]
        if paras:
            grids.append(Grid("Texte", "line", _norm_rows([[p] for p in paras])))
    return Loaded(kind, grids)


# --------------------------------------------------------------------------------------------------- CSV / texte / Markdown
def _split_text_line(line: str) -> List[str]:
    if "\t" in line:
        cells = line.split("\t")
    elif "|" in line:
        cells = line.split("|")
        if cells and not cells[0].strip():
            cells = cells[1:]
        if cells and not cells[-1].strip():
            cells = cells[:-1]
    elif re.search(r"\s;\s|;\s*\S.*;", line):
        cells = line.split(";")
    else:
        cells = re.split(r"\s{2,}", line.strip())
    return [clean_spaces(c) for c in cells]


def _read_csv(data: bytes) -> Loaded:
    text = _decode(data)
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") >= sample.count(",") else ","
    rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return Loaded("csv", [Grid("Fichier", "line", _norm_rows(rows))])


def _read_txt(data: bytes) -> Loaded:
    rows = []
    for line in _decode(data).splitlines():
        if not line.strip():
            rows.append([""])
            continue
        cells = _split_text_line(line)
        if cells and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):   # séparateur de tableau Markdown
            rows.append([""])
            continue
        rows.append(cells)
    return Loaded("txt", [Grid("Fichier", "line", _norm_rows(rows))])


# --------------------------------------------------------------------------------------------------- HTML
class _H(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "header", "footer", "ul", "ol"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: List[List[List[str]]] = []
        self.stack: List[dict] = []
        self.lines: List[str] = []
        self.buf: List[str] = []
        self.skip = 0

    def _flush(self):
        t = clean_spaces("".join(self.buf))
        self.buf = []
        if t:
            self.lines.append(t)

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
        elif tag == "table":
            if not self.stack:
                self._flush()
            self.stack.append({"rows": [], "row": None, "cell": None})
        elif self.stack:
            s = self.stack[-1]
            if tag == "tr":
                s["row"] = []
            elif tag in ("td", "th"):
                s["cell"] = []
                span = dict(attrs).get("colspan")
                s["span"] = int(span) if span and span.isdigit() else 1
        elif tag in self.BLOCK:
            self._flush()

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip = max(0, self.skip - 1)
        elif tag == "table" and self.stack:
            s = self.stack.pop()
            if s["row"]:
                s["rows"].append(s["row"])
            if s["rows"] and len(self.tables) < MAX_GRIDS:
                self.tables.append(s["rows"])
        elif self.stack:
            s = self.stack[-1]
            if tag in ("td", "th") and s["cell"] is not None and s["row"] is not None:
                s["row"].append(clean_spaces("".join(s["cell"])))
                s["row"].extend([""] * (s.get("span", 1) - 1))
                s["cell"] = None
            elif tag == "tr" and s["row"] is not None:
                s["rows"].append(s["row"])
                s["row"] = None
        elif tag in self.BLOCK:
            self._flush()

    def handle_data(self, data):
        if self.skip:
            return
        if self.stack:
            if self.stack[-1]["cell"] is not None:
                self.stack[-1]["cell"].append(data)
        else:
            self.buf.append(data)


def _read_html(data: bytes) -> Loaded:
    p = _H()
    p.feed(_decode(data))
    p._flush()
    grids = [Grid(f"Tableau {i + 1}", "table", _norm_rows(t)) for i, t in enumerate(p.tables)]
    grids = [g for g in grids if g.rows]
    if p.lines:
        grids.append(Grid("Texte", "line", _norm_rows([[ln] for ln in p.lines])))
    return Loaded("html", grids)


# --------------------------------------------------------------------------------------------------- JSON
def _read_json(data: bytes) -> Loaded:
    obj = _try_json(_decode(data))
    if obj is None:
        raise UnsupportedFormat("JSON invalide.")
    found: List[Tuple[list, dict]] = []

    def walk(o, parent):
        if isinstance(o, list) and o and all(isinstance(x, dict) for x in o):
            found.append((o, parent))
            return
        if isinstance(o, dict):
            for v in o.values():
                walk(v, o)
        elif isinstance(o, list):
            for v in o:
                walk(v, parent)

    walk(obj, obj if isinstance(obj, dict) else {})
    grids: List[Grid] = []
    if found:
        items, _ = max(found, key=lambda f: len(f[0]))
        keys: List[str] = []
        for it in items:
            for k in it:
                if k not in keys and not isinstance(it[k], (dict, list)):
                    keys.append(k)
        rows = [keys] + [[_num_text(it.get(k)) if not isinstance(it.get(k), bool) else str(it.get(k)) for k in keys] for it in items]
        grids.append(Grid("Éléments", "item", _norm_rows(rows)))

    fields: List[List[str]] = []

    def collect(o, depth=0):
        if isinstance(o, dict) and depth < 4:
            for k, v in o.items():
                if isinstance(v, dict):
                    collect(v, depth + 1)
                elif not isinstance(v, list):
                    fields.append([str(k), _num_text(v)])

    collect(obj)
    if fields:
        grids.append(Grid("Champs", "field", _norm_rows(fields)))
    return Loaded("json", grids)


def load(data: bytes, filename: str = "") -> Loaded:
    kind = detect_kind(data, filename)
    if kind == "xlsx":
        return _read_xlsx(data)
    if kind == "docx":
        return _read_docx(data)
    if kind in ("ods", "odt"):
        return _read_odf(data, kind)
    if kind == "csv":
        return _read_csv(data)
    if kind == "html":
        return _read_html(data)
    if kind == "json":
        return _read_json(data)
    if kind == "image":
        return Loaded("image", [], warnings=["Image : document scanné ou photographié. Hors périmètre (pas d'OCR)."])
    return _read_txt(data)
