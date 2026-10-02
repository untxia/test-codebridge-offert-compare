#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Mêmes offres (original S0 / révision S1) que les PDF, écrites dans d'autres formats : xlsx, docx, ods, csv, html, md, txt, json.
Sert à vérifier que l'outil lit « n'importe quel format » et trouve les MÊMES écarts, y compris en mélangeant les formats.
Sorties : testset/formats/{original,revised}.<ext>.  Dépendances de génération : openpyxl, python-docx, odfpy."""
import csv
import datetime as dt
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import generate_testset as g  # noqa: E402

OUT = os.path.join(HERE, "formats")
D = g.D


def rows_of(name):
    o = g.OFFERS[name]
    return o, o["rows"], g.stated_totals(o)


def qty_txt(r):
    return f"{int(r.qty)}" + (f" {r.unit}" if r.unit else "")


def table_text(name):
    """Grille texte avec le style de mise en forme du PDF correspondant."""
    o, rows, tot = rows_of(name)
    st = o["style"]
    head = st["columns"]
    body = [[str(i), r.label, qty_txt(r), st["money"](r.price), st["money"](r.total), st["date"](r.delivery)] for i, r in enumerate(rows, 1)]
    lab = st["totals_labels"]
    totals = [(lab[k], st["money"](tot[k])) for k in ("ht", "tva", "ttc")]
    return o, rows, head, body, totals


def write_xlsx(name, path):
    import openpyxl
    from openpyxl.styles import Font
    o, rows, tot = rows_of(name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Offre"
    ws["A1"], ws["A2"], ws["A3"] = f"Offre commerciale {g.OFFER_REF}", f"Fournisseur : {g.SUPPLIER}", f"Client : {g.CLIENT}"
    ws["A1"].font = Font(bold=True, size=14)
    heads = o["style"]["columns"]
    for c, h in enumerate(heads, 1):
        ws.cell(5, c, h).font = Font(bold=True)
    for i, r in enumerate(rows, 1):
        row = 5 + i
        ws.cell(row, 1, i)
        ws.cell(row, 2, r.label)
        ws.cell(row, 3, f"{int(r.qty)} {r.unit}" if r.unit else int(r.qty))
        ws.cell(row, 4, float(r.price)).number_format = '#,##0.00 "€"'
        ws.cell(row, 5, float(r.total)).number_format = '#,##0.00 "€"'
        if len(r.delivery) == 10 and r.delivery[4] == "-":
            ws.cell(row, 6, dt.date.fromisoformat(r.delivery)).number_format = "DD/MM/YYYY"
        else:
            ws.cell(row, 6, r.delivery)
    base = 5 + len(rows) + 2
    for k, (lab, key) in enumerate((("Total HT", "ht"), ("TVA 20 %", "tva"), ("Total TTC", "ttc"))):
        ws.cell(base + k, 5, lab)
        ws.cell(base + k, 6, float(tot[key])).number_format = '#,##0.00 "€"'
    ws2 = wb.create_sheet("Conditions")
    ws2["A1"] = g.CONDITIONS
    wb.save(path)


def write_docx(name, path):
    import docx
    o, rows, head, body, totals = table_text(name)
    d = docx.Document()
    d.add_heading(o["style"]["title"], 1)
    for line in (f"Fournisseur : {g.SUPPLIER}", f"Client : {g.CLIENT}", "Devise : EUR"):
        d.add_paragraph(line)
    t = d.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        t.rows[0].cells[i].text = h
    for row in body:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    d.add_paragraph("")
    for lab, val in totals:
        d.add_paragraph(f"{lab} : {val}")
    d.add_paragraph(g.CONDITIONS)
    d.save(path)


def write_ods(name, path):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell
    from odf.text import P
    o, rows, tot = rows_of(name)
    doc = OpenDocumentSpreadsheet()
    t = Table(name="Offre")

    def cell(v):
        if isinstance(v, (int, float)):
            c = TableCell(valuetype="float", value=str(v))
        else:
            c = TableCell(valuetype="string")
        c.addElement(P(text=str(v)))
        return c

    def row(vals):
        r = TableRow()
        for v in vals:
            r.addElement(cell(v))
        t.addElement(r)

    row([f"Offre commerciale {g.OFFER_REF}"])
    row([])
    row(o["style"]["columns"])
    for i, r in enumerate(rows, 1):
        row([i, r.label, qty_txt(r), float(r.price), float(r.total), g.date_a(r.delivery)])
    row([])
    row(["", "", "", "Total HT", float(tot["ht"])])
    row(["", "", "", "TVA 20 %", float(tot["tva"])])
    row(["", "", "", "Total TTC", float(tot["ttc"])])
    doc.spreadsheet.addElement(t)
    doc.save(path)


def write_csv(name, path):
    o, rows, head, body, totals = table_text(name)
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow([f"Offre {g.OFFER_REF}"])
        w.writerow(head)
        w.writerows(body)
        w.writerow([])
        for lab, val in totals:
            w.writerow(["", "", "", lab, val])


def write_html(name, path):
    import html
    o, rows, head, body, totals = table_text(name)
    h = html.escape
    s = [f"<!doctype html><html><body><h1>{h(o['style']['title'])}</h1><p>Fournisseur : {h(g.SUPPLIER)}</p><p>Devise : EUR</p><table>"]
    s.append("<tr>" + "".join(f"<th>{h(c)}</th>" for c in head) + "</tr>")
    for row in body:
        s.append("<tr>" + "".join(f"<td>{h(c)}</td>" for c in row) + "</tr>")
    s.append("</table>")
    for lab, val in totals:
        s.append(f"<p><b>{h(lab)}</b> : {h(val)}</p>")
    s.append("</body></html>")
    open(path, "w", encoding="utf-8").write("\n".join(s))


def write_md(name, path):
    o, rows, head, body, totals = table_text(name)
    lines = [f"# {o['style']['title']}", "", "| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    lines += [""] + [f"**{lab}** : {val}" for lab, val in totals]
    open(path, "w", encoding="utf-8").write("\n".join(lines) + "\n")


def write_txt(name, path):
    o, rows, head, body, totals = table_text(name)
    widths = [max(len(r[i]) for r in [head] + body) + 3 for i in range(len(head))]
    fmt = lambda r: "".join(c.ljust(w) for c, w in zip(r, widths)).rstrip()  # noqa: E731
    lines = [o["style"]["title"], f"Devise : EUR", "", fmt(head)] + [fmt(r) for r in body] + [""]
    lines += [f"{lab}   {val}" for lab, val in totals]
    open(path, "w", encoding="utf-8").write("\n".join(lines) + "\n")


def write_json(name, path):
    o, rows, tot = rows_of(name)
    doc = {"reference": g.OFFER_REF, "currency": "EUR",
           "lines": [{"no": i, "description": r.label, "quantity": qty_txt(r), "unit_price": float(r.price),
                      "line_total": float(r.total), "delivery_date": r.delivery} for i, r in enumerate(rows, 1)],
           "totals": {"total_ht": float(tot["ht"]), "vat": float(tot["tva"]), "total_ttc": float(tot["ttc"])}}
    json.dump(doc, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=2)


WRITERS = {"xlsx": write_xlsx, "docx": write_docx, "ods": write_ods, "csv": write_csv, "html": write_html,
           "md": write_md, "txt": write_txt, "json": write_json}


def main():
    os.makedirs(OUT, exist_ok=True)
    for ext, fn in WRITERS.items():
        for name in ("original", "revised"):
            fn(name, os.path.join(OUT, f"{name}.{ext}"))
    print("OK ->", OUT, sorted(os.listdir(OUT)))


if __name__ == "__main__":
    main()
